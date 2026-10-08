"""Модуль экспорта отчётов в Excel (openpyxl) и печатных форм (HTML / PDF-ready).

Соответствует разделу 7 кейса («Отчёты и аналитика», «Выгрузка в PDF / Excel»).
Полностью автономен, без внешних API и облачных сервисов.
"""
from __future__ import annotations

import html
import io
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from ..models import (
    AIAssessment,
    Employee,
    Equipment,
    FaultCode,
    Material,
    MaterialNorm,
    MaterialWriteOff,
    Role,
    Section,
    Status,
    WorkOrder,
    WorkOrderEvent,
)
from ..serializers import (
    PRIORITY_LABELS,
    STATUS_LABELS,
    downtime_minutes,
    short_name,
    work_minutes,
)
from ..i18n import T, is_kz
from . import reports

# Русские подписи с сервера (статусы, приоритеты, участки, шифры, специальности) → казахский
_KZ_LABELS = {
    "Выдан": "Берілді", "В очереди": "Кезекте", "Принят": "Қабылданды", "Отклонён": "Бас тартылды",
    "В работе": "Жұмыста", "Приостановлен": "Тоқтатылды", "Исполнено": "Орындалды",
    "Проверка ИИ": "ЖИ тексеруі", "На доработке": "Қайта қарауда", "Закрыт": "Жабылды", "Отменён": "Болдырылмады",
    "Аварийный": "Апаттық", "Высокий": "Жоғары", "Обычный": "Қалыпты", "Плановый": "Жоспарлы",
    "Принято": "Қабылданды", "Принято с замечаниями": "Ескертулермен қабылданды",
    "Требует доработки": "Қайта қарауды қажет етеді",
    "Слесарь": "Слесарь", "Электрик": "Электрик", "Сварщик": "Дәнекерлеуші",
    "Мастер смены": "Ауысым шебері", "Главный механик": "Бас механик", "Администратор": "Әкімші",
    "Участок дробления": "Ұсақтау бөлімшесі", "Обогатительная фабрика": "Байыту фабрикасы",
    "Участок сушки": "Кептіру бөлімшесі", "Ремонтно-механический цех": "Жөндеу-механикалық цех",
    "Износ футеровки / бронеплит": "Футеровка / сауыт тақталарының тозуы",
    "Разрушение подшипника": "Мойынтіректің бұзылуы",
    "Износ роликов / барабана": "Ролик / барабанның тозуы",
    "Порыв конвейерной ленты": "Конвейер лентасының үзілуі",
    "Неисправность редуктора": "Редуктордың ақауы",
    "Ослабление крепежа, вибрация": "Бекітпенің босауы, діріл",
    "Трещина металлоконструкции": "Металл құрылымдағы жарықшақ",
    "Износ зубчатой передачи": "Тісті берілістің тозуы",
    "Отказ электродвигателя": "Электр қозғалтқыштың істен шығуы",
    "Повреждение кабеля": "Кабельдің зақымдануы",
    "Неисправность пускателя / автомата": "Іске қосқыш / автоматтың ақауы",
    "Отказ датчика": "Датчиктің істен шығуы",
    "Неисправность освещения": "Жарықтандырудың ақауы",
    "Течь масла / гидрожидкости": "Май / гидрсұйықтықтың ағуы",
    "Отказ гидронасоса": "Гидрсорғының істен шығуы",
    "Повреждение РВД": "Жоғары қысымды шлангының зақымдануы",
    "Утечка сжатого воздуха": "Сығылған ауаның ағуы",
    "Неисправность пневмоцилиндра": "Пневмоцилиндрдің ақауы",
    "Нарушение смазки узла": "Тораптың майлануының бұзылуы",
    "Загрязнение / замена масла": "Майдың ластануы / ауыстыру",
}


def _kz(text):
    """Подпись из справочника на языке запроса (для неизвестных текстов — как есть)."""
    if text is None or not is_kz():
        return text
    return _KZ_LABELS.get(text, text)

# Цветовая палитра оформления АО «Костанайские Минералы»
COLOR_HEADER_BG = "1E293B"      # slate-800
COLOR_HEADER_FG = "FFFFFF"
COLOR_ACCENT = "059669"         # emerald-600
COLOR_SUBHEADER_BG = "F1F5F9"   # slate-100
COLOR_BORDER = "CBD5E1"         # slate-300
COLOR_ALERT_BG = "FEE2E2"       # red-100
COLOR_ALERT_FG = "991B1B"       # red-800


def _style_header_row(ws, row_idx: int, col_count: int, bg_hex: str = COLOR_HEADER_BG, fg_hex: str = COLOR_HEADER_FG):
    fill = PatternFill(start_color=bg_hex, end_color=bg_hex, fill_type="solid")
    font = Font(name="Arial", size=10, bold=True, color=fg_hex)
    thin = Side(border_style="thin", color=COLOR_BORDER)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col in range(1, col_count + 1):
        cell = ws.cell(row=row_idx, column=col)
        cell.fill = fill
        cell.font = font
        cell.border = border
        cell.alignment = align
    ws.row_dimensions[row_idx].height = 28


def _apply_table_borders(ws, start_row: int, end_row: int, col_count: int):
    thin = Side(border_style="thin", color="E2E8F0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for r in range(start_row, end_row + 1):
        for c in range(1, col_count + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = border
            if cell.font.name != "Arial":
                cell.font = Font(name="Arial", size=9)


def _auto_column_widths(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or "")
            if len(val) > max_len and len(val) < 80:
                max_len = len(val)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 11)


# =====================================================================
# 1. Экспорт отчёта за смену в Excel
# =====================================================================
def export_shift_report_excel(db: Session, start: datetime, end: datetime,
                              section_id: int | None = None, brigade_id: int | None = None) -> io.BytesIO:
    data = reports.shift_report(db, start, end, section_id, brigade_id)
    wb = Workbook()

    # Лист 1: Сводка смены
    ws_sum = wb.active
    ws_sum.title = T("Сводка смены", "Ауысым қорытындысы")
    ws_sum.views.sheetView[0].showGridLines = True

    # Заголовок документа
    ws_sum.merge_cells("A1:G1")
    title_cell = ws_sum["A1"]
    title_cell.value = T("АО «Костанайские Минералы» — Сводный отчёт за смену", "«Қостанай минералдары» АҚ — Ауысымның жиынтық есебі")
    title_cell.font = Font(name="Arial", size=14, bold=True, color="065F46")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_sum.row_dimensions[1].height = 35

    ws_sum["A2"] = f"{T('Период', 'Кезең')}: {start.strftime('%d.%m.%Y %H:%M')} — {end.strftime('%d.%m.%Y %H:%M')}"
    ws_sum["A2"].font = Font(name="Arial", size=10, italic=True)

    # Таблица показателей
    metrics = [
        (T("Выдано нарядов за смену", "Ауысымда берілген нарядтар"), data["issued"]),
        (T("Выполнено нарядов (завершено)", "Орындалған нарядтар (аяқталған)"), data["done"]),
        (T("Закрыто и утверждено мастером", "Шебер жауып, бекіткен"), data["closed"]),
        (T("Просрочено нарядов", "Мерзімі өткен нарядтар"), data["overdue"]),
        (T("Отклонений исполнителями", "Орындаушылардың бас тартулары"), data["rejected"]),
        (T("Суммарный простой оборудования (часов)", "Жабдықтың жиынтық тоқтап тұруы (сағат)"), f"{data['downtime_hours']} {T('ч', 'сағ')}"),
        (T("Средняя оценка качества закрытия (ИИ)", "Жабудың орташа сапа бағасы (ЖИ)"), f"{data['avg_score']}/100" if data['avg_score'] else "—"),
    ]

    ws_sum.cell(row=4, column=1, value=T("Показатель смены", "Ауысым көрсеткіші"))
    ws_sum.cell(row=4, column=2, value=T("Значение", "Мәні"))
    _style_header_row(ws_sum, 4, 2, bg_hex="065F46", fg_hex="FFFFFF")

    cur_row = 5
    for title, val in metrics:
        ws_sum.cell(row=cur_row, column=1, value=title).font = Font(name="Arial", size=10)
        c_val = ws_sum.cell(row=cur_row, column=2, value=val)
        c_val.font = Font(name="Arial", size=10, bold=True)
        c_val.alignment = Alignment(horizontal="center")
        cur_row += 1
    _apply_table_borders(ws_sum, 5, cur_row - 1, 2)

    # ИИ-резюме
    cur_row += 1
    ws_sum.cell(row=cur_row, column=1, value=T("Итоговое ИИ-резюме смены:", "Ауысымның қорытынды ЖИ-түйіндемесі:")).font = Font(name="Arial", size=10, bold=True)
    cur_row += 1
    ws_sum.merge_cells(start_row=cur_row, start_column=1, end_row=cur_row + 2, end_column=6)
    summary_cell = ws_sum.cell(row=cur_row, column=1, value=data["summary"])
    summary_cell.font = Font(name="Arial", size=9, italic=True)
    summary_cell.alignment = Alignment(wrap_text=True, vertical="top")
    summary_cell.fill = PatternFill(start_color="F0FDF4", end_color="F0FDF4", fill_type="solid")

    # Таблица загрузки сотрудников
    cur_row += 4
    ws_sum.cell(row=cur_row, column=1, value=T("Загрузка ремонтно-технического персонала", "Жөндеу-техникалық персоналдың жүктемесі")).font = Font(name="Arial", size=11, bold=True)
    cur_row += 1
    load_headers = [T("Сотрудник", "Қызметкер"), T("Нарядов", "Нарядтар"), T("Отработано (мин)", "Жұмыс істеді (мин)"), T("Отработано (ч)", "Жұмыс істеді (сағ)")]
    for c_idx, h in enumerate(load_headers, 1):
        ws_sum.cell(row=cur_row, column=c_idx, value=h)
    _style_header_row(ws_sum, cur_row, len(load_headers))
    load_start = cur_row + 1
    cur_row += 1
    for w in data.get("load", []):
        ws_sum.cell(row=cur_row, column=1, value=w["name"])
        ws_sum.cell(row=cur_row, column=2, value=w["orders"]).alignment = Alignment(horizontal="center")
        ws_sum.cell(row=cur_row, column=3, value=w["minutes"]).alignment = Alignment(horizontal="center")
        ws_sum.cell(row=cur_row, column=4, value=round(w["minutes"] / 60, 1)).alignment = Alignment(horizontal="center")
        cur_row += 1
    if data.get("load"):
        _apply_table_borders(ws_sum, load_start, cur_row - 1, len(load_headers))

    _auto_column_widths(ws_sum)

    # Лист 2: Реестр нарядов смены
    ws_orders = wb.create_sheet(title=T("Наряды смены", "Ауысым нарядтары"))
    ws_orders.views.sheetView[0].showGridLines = True
    order_headers = [
        T("№ наряда", "Наряд №"), T("Оборудование", "Жабдық"), T("Участок", "Бөлімше"),
        T("Приоритет", "Басымдық"), T("Тип", "Түрі"),
        T("Исполнитель", "Орындаушы"), T("Статус", "Мәртебесі"), T("Срок", "Мерзімі"),
        T("Просрочен", "Мерзімі өтті"), T("Время (мин)", "Уақыт (мин)"),
        T("Оценка ИИ", "ЖИ бағасы"), T("Шифр дефекта", "Ақау шифры"),
    ]
    for c_idx, h in enumerate(order_headers, 1):
        ws_orders.cell(row=1, column=c_idx, value=h)
    _style_header_row(ws_orders, 1, len(order_headers))

    orders = reports._period_orders(db, start, end, section_id, brigade_id)
    r_idx = 2
    for o in orders:
        ws_orders.cell(row=r_idx, column=1, value=o.number).alignment = Alignment(horizontal="center")
        ws_orders.cell(row=r_idx, column=2, value=o.equipment.name if o.equipment else "—")
        ws_orders.cell(row=r_idx, column=3, value=_kz(o.section.name) if o.section else "—")
        ws_orders.cell(row=r_idx, column=4, value=_kz(PRIORITY_LABELS.get(o.priority, str(o.priority))))
        ws_orders.cell(row=r_idx, column=5, value=T("Внеплановый", "Жоспардан тыс") if o.work_type == "unplanned" else T("Плановый", "Жоспарлы"))
        ws_orders.cell(row=r_idx, column=6, value=short_name(o.assignee.full_name) if o.assignee else "—")
        ws_orders.cell(row=r_idx, column=7, value=_kz(STATUS_LABELS.get(Status(o.status), str(o.status))))
        ws_orders.cell(row=r_idx, column=8, value=o.deadline.strftime("%H:%M %d.%m") if o.deadline else "—")
        ws_orders.cell(row=r_idx, column=9, value=T("Да", "Иә") if reports.is_overdue(o) else T("Нет", "Жоқ")).alignment = Alignment(horizontal="center")
        minutes = work_minutes(o)
        ws_orders.cell(row=r_idx, column=10, value=int(minutes) if minutes is not None else "—").alignment = Alignment(horizontal="center")
        score = o.assessment.final_score if o.assessment else None
        ws_orders.cell(row=r_idx, column=11, value=score if score is not None else "—").alignment = Alignment(horizontal="center")
        ws_orders.cell(row=r_idx, column=12, value=o.fault_code.code if o.fault_code else "—").alignment = Alignment(horizontal="center")
        r_idx += 1

    if len(orders) > 0:
        _apply_table_borders(ws_orders, 2, r_idx - 1, len(order_headers))
    _auto_column_widths(ws_orders)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


# =====================================================================
# 2. Экспорт рейтинга рабочих и бригад в Excel
# =====================================================================
def export_rating_excel(db: Session, start: datetime, end: datetime, brigade_id: int | None = None) -> io.BytesIO:
    rows = reports.compute_rating(db, start, end, brigade_id)
    wb = Workbook()
    ws = wb.active
    ws.title = T("Рейтинг сотрудников", "Қызметкерлер рейтингі")
    ws.views.sheetView[0].showGridLines = True

    # Заголовок
    ws.merge_cells("A1:K1")
    t = ws["A1"]
    t.value = T("АО «Костанайские Минералы» — Объективный рейтинг ремонтно-технического персонала", "«Қостанай минералдары» АҚ — Жөндеу-техникалық персоналдың объективті рейтингі")
    t.font = Font(name="Arial", size=13, bold=True, color="065F46")
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32

    ws["A2"] = (
        f"{T('Период', 'Кезең')}: {start.strftime('%d.%m.%Y')} — {end.strftime('%d.%m.%Y')} | "
        f"{T('Формула', 'Формула')}: " + T(
            "0.35·Качество + 0.25·В срок + 0.20·(100 − Повторы) + 0.15·Объём + 0.05·(100 − Отказы)",
            "0.35·Сапа + 0.25·Мерзімінде + 0.20·(100 − Қайталаулар) + 0.15·Көлем + 0.05·(100 − Бас тартулар)")
    )
    ws["A2"].font = Font(name="Arial", size=9, italic=True)

    headers = [
        T("Место", "Орны"), T("Сотрудник", "Қызметкер"), T("Специальность", "Мамандығы"),
        T("Разряд", "Разряд"), T("Бригада", "Бригада"),
        T("Итоговый балл", "Қорытынды балл"), T("Качество (35%)", "Сапа (35%)"),
        T("В срок (25%)", "Мерзімінде (25%)"), T("Без повторов (20%)", "Қайталаусыз (20%)"),
        T("Объём/сложность (15%)", "Көлем/күрделілік (15%)"), T("Без отказов (5%)", "Бас тартусыз (5%)"),
        T("Закрыто нарядов", "Жабылған нарядтар"), T("Повторов за 7 дней", "7 күндегі қайталаулар"),
        T("ИИ-рекомендация исполнителю", "Орындаушыға ЖИ ұсынысы"),
    ]
    for c_idx, h in enumerate(headers, 1):
        ws.cell(row=4, column=c_idx, value=h)
    _style_header_row(ws, 4, len(headers), bg_hex="065F46", fg_hex="FFFFFF")

    r_idx = 5
    for row in rows:
        place_cell = ws.cell(row=r_idx, column=1, value=row.get("place", r_idx - 4))
        place_cell.alignment = Alignment(horizontal="center")
        place_cell.font = Font(name="Arial", size=10, bold=True)
        ws.cell(row=r_idx, column=2, value=row["full_name"]).font = Font(name="Arial", size=10, bold=True)
        ws.cell(row=r_idx, column=3, value=_kz(row["specialty"]))
        ws.cell(row=r_idx, column=4, value=row["grade"]).alignment = Alignment(horizontal="center")
        ws.cell(row=r_idx, column=5, value=row.get("brigade", {}).get("name", "—") if row.get("brigade") else "—")
        
        # Рейтинг
        rate_cell = ws.cell(row=r_idx, column=6, value=row["rating"])
        rate_cell.alignment = Alignment(horizontal="center")
        rate_cell.font = Font(name="Arial", size=11, bold=True, color="065F46")
        rate_cell.fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")

        comp = row.get("components", {})
        ws.cell(row=r_idx, column=7, value=comp.get("quality", 0)).alignment = Alignment(horizontal="center")
        ws.cell(row=r_idx, column=8, value=comp.get("on_time", 0)).alignment = Alignment(horizontal="center")
        ws.cell(row=r_idx, column=9, value=comp.get("no_rework", 0)).alignment = Alignment(horizontal="center")
        ws.cell(row=r_idx, column=10, value=comp.get("volume", 0)).alignment = Alignment(horizontal="center")
        ws.cell(row=r_idx, column=11, value=comp.get("no_reject", 0)).alignment = Alignment(horizontal="center")

        ws.cell(row=r_idx, column=12, value=row["orders_closed"]).alignment = Alignment(horizontal="center")
        ws.cell(row=r_idx, column=13, value=row["repeat_failures"]).alignment = Alignment(horizontal="center")
        ws.cell(row=r_idx, column=14, value=row.get("explanation", ""))
        r_idx += 1

    if rows:
        _apply_table_borders(ws, 5, r_idx - 1, len(headers))
    _auto_column_widths(ws)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


# =====================================================================
# 3. Экспорт списанных ТМЦ и отклонений от норм в Excel
# =====================================================================
def export_materials_excel(db: Session, start: datetime, end: datetime) -> io.BytesIO:
    materials_stat = get_materials_report(db, start, end)
    wb = Workbook()
    ws = wb.active
    ws.title = T("Списание ТМЦ", "ТМҚ есептен шығару")
    ws.views.sheetView[0].showGridLines = True

    ws.merge_cells("A1:G1")
    t = ws["A1"]
    t.value = T("АО «Костанайские Минералы» — Отчёт по списанию запчастей и ТМЦ против нормативов", "«Қостанай минералдары» АҚ — Қосалқы бөлшектер мен ТМҚ-ны нормативтерге қарсы есептен шығару есебі")
    t.font = Font(name="Arial", size=13, bold=True, color="1E3A8A")
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32

    headers = [
        T("Наименование ТМЦ / запчасти", "ТМҚ / қосалқы бөлшек атауы"), T("Ед. изм.", "Өлшем бірлігі"),
        T("Фактически списано", "Нақты есептен шығарылды"),
        T("Типичный норматив", "Типтік норматив"), T("Отклонение от нормы", "Нормадан ауытқу"),
        T("Кол-во нарядов", "Нарядтар саны"), T("Статус контроля ИИ", "ЖИ бақылау мәртебесі"),
    ]
    for c_idx, h in enumerate(headers, 1):
        ws.cell(row=3, column=c_idx, value=h)
    _style_header_row(ws, 3, len(headers), bg_hex="1E3A8A", fg_hex="FFFFFF")

    r_idx = 4
    for item in materials_stat["items"]:
        ws.cell(row=r_idx, column=1, value=item["name"]).font = Font(name="Arial", size=10, bold=True)
        ws.cell(row=r_idx, column=2, value=item["unit"]).alignment = Alignment(horizontal="center")
        ws.cell(row=r_idx, column=3, value=item["total_qty"]).alignment = Alignment(horizontal="center")
        ws.cell(row=r_idx, column=4, value=item["norm_qty"]).alignment = Alignment(horizontal="center")

        diff = item.get("diff_pct", 0)
        c_diff = ws.cell(row=r_idx, column=5, value=f"{diff:+.1f}%")
        c_diff.alignment = Alignment(horizontal="center")
        if diff > 50:
            c_diff.font = Font(name="Arial", size=10, bold=True, color="DC2626")
            c_diff.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")

        ws.cell(row=r_idx, column=6, value=item["orders_count"]).alignment = Alignment(horizontal="center")
        ws.cell(row=r_idx, column=7, value=item["status_label"])
        r_idx += 1

    if materials_stat["items"]:
        _apply_table_borders(ws, 4, r_idx - 1, len(headers))
    _auto_column_widths(ws)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


# =====================================================================
# 4. JSON-агрегация по списанным материалам (для UI и экспорта)
# =====================================================================
def get_materials_report(db: Session, start: datetime, end: datetime) -> dict:
    orders = reports._period_orders(db, start, end)
    order_ids = [o.id for o in orders]
    if not order_ids:
        return {"items": [], "total_writeoffs": 0, "anomalies_count": 0}

    writeoffs = db.scalars(
        select(MaterialWriteOff)
        .where(MaterialWriteOff.order_id.in_(order_ids))
        .options(selectinload(MaterialWriteOff.material), selectinload(MaterialWriteOff.order))
    ).all()

    # Справочник норм
    norms = db.scalars(select(MaterialNorm)).all()
    norm_map = {(n.fault_code_id, n.material_id): n.typical_qty for n in norms}

    by_mat: dict[int, dict] = {}
    for w in writeoffs:
        mid = w.material_id
        if mid not in by_mat:
            mat_name = w.material.name if w.material else f"{T('Материал', 'Материал')} #{mid}"
            mat_unit = w.material.unit if w.material else T("ед.", "бірл.")
            by_mat[mid] = {
                "material_id": mid,
                "name": mat_name,
                "unit": mat_unit,
                "total_qty": 0.0,
                "expected_qty": 0.0,
                "orders": set(),
                "overuse_events": 0,
            }
        by_mat[mid]["total_qty"] += w.qty
        by_mat[mid]["orders"].add(w.order_id)
        if w.order.fault_code_id:
            typ = norm_map.get((w.order.fault_code_id, mid), w.qty)
            by_mat[mid]["expected_qty"] += typ
            if w.qty > typ * 1.5:
                by_mat[mid]["overuse_events"] += 1
        else:
            by_mat[mid]["expected_qty"] += w.qty

    items = []
    anomalies_count = 0
    for stat in by_mat.values():
        total_q = round(stat["total_qty"], 1)
        expected_q = round(stat["expected_qty"], 1)
        diff_pct = round(((total_q - expected_q) / expected_q * 100), 1) if expected_q > 0 else 0.0
        is_anom = diff_pct > 40 and stat["overuse_events"] >= 2
        if is_anom:
            anomalies_count += 1
        items.append({
            "material_id": stat["material_id"],
            "name": stat["name"],
            "unit": stat["unit"],
            "total_qty": total_q,
            "norm_qty": expected_q,
            "diff_pct": diff_pct,
            "orders_count": len(stat["orders"]),
            "overuse_count": stat["overuse_events"],
            "is_anomaly": is_anom,
            "status_label": T("⚠️ Перерасход выше нормы", "⚠️ Норма үстіндегі артық шығын") if is_anom else T("Нормативный расход", "Нормативтік шығын"),
        })

    items.sort(key=lambda x: (x["is_anomaly"], x["diff_pct"]), reverse=True)
    return {
        "period": {"start": start, "end": end},
        "items": items,
        "total_writeoffs": len(writeoffs),
        "anomalies_count": anomalies_count,
    }


# =====================================================================
# 5. Печатная форма наряда (Print-Ready / PDF)
# =====================================================================
def _e(value) -> str:
    """Экранирование значений из БД перед вставкой в HTML (защита от XSS)."""
    return html.escape(str(value), quote=True)


def generate_order_print_html(order: WorkOrder) -> str:
    """Генерирует официальную форму наряда-допуска АО «Костанайские Минералы».

    Все значения из БД (тексты исполнителя/мастера, справочники) экранируются через _e():
    страница открывается с того же origin, что и SPA, где хранится JWT.
    """
    events_html = "".join(
        f"<tr><td>{ev.created_at.strftime('%d.%m.%Y %H:%M:%S') if ev.created_at else '—'}</td>"
        f"<td><strong>{_e(ev.action)}</strong></td>"
        f"<td>{_e(short_name(ev.actor.full_name)) if ev.actor else T('ИИ / Система', 'ЖИ / Жүйе')}</td>"
        f"<td>{_e(ev.comment or ev.reason or '—')}</td></tr>"
        for ev in (order.events or [])
    )

    mats_html = "".join(
        f"<tr><td>{_e(m.material.name) if m.material else f"{T('Материал', 'Материал')} #{m.material_id}"}</td>"
        f"<td>{_e(m.qty)} {_e(m.material.unit) if m.material else T('ед.', 'бірл.')}</td></tr>"
        for m in (order.materials or [])
    ) or "<tr><td colspan='2' style='text-align:center; color:#64748b;'>{}</td></tr>".format(T('Материалы не списывались', 'Материалдар есептен шығарылмады'))

    eq_desc = (f"{_e(order.equipment.name)} ({T('Инв. №', 'Инв. №')} {_e(order.equipment.inv_no)})" if order.equipment
               else T("Не указано (Инв. № —)", "Көрсетілмеген (Инв. № —)"))
    sec_name = _e(_kz(order.section.name)) if order.section else T("Не указан", "Көрсетілмеген")
    created_str = order.created_at.strftime('%d.%m.%Y %H:%M') if order.created_at else "—"
    deadline_str = order.deadline.strftime('%d.%m.%Y %H:%M') if order.deadline else "—"
    master_name = _e(order.master.full_name) if order.master else T("Мастер смены", "Ауысым шебері")
    assignee_name = _e(order.assignee.full_name) if order.assignee else T("Не назначен", "Тағайындалмаған")
    assignee_spec = _e(_kz(order.assignee.specialty)) if order.assignee else "—"
    priority_label = _e(_kz(PRIORITY_LABELS.get(order.priority, str(order.priority))))
    description = _e(order.description or '—')
    comment = _e(order.comment) if order.comment else ''
    fault_code = _e(order.fault_code.code) if order.fault_code else '—'
    fault_name = _e(_kz(order.fault_code.name)) if order.fault_code else '—'
    work_done = _e(order.work_done) if order.work_done else T('Работы не описаны', 'Жұмыстар сипатталмаған')
    close_comment = _e(order.close_comment) if order.close_comment else ''

    assessment_html = ""
    if order.assessment:
        a = order.assessment
        v_raw = a.verdict.value if hasattr(a.verdict, "value") else str(a.verdict)
        verdict_trans = _e({
            "accepted": T("✅ Принято без замечаний", "✅ Ескертусіз қабылданды"),
            "accepted_with_remarks": T("⚠️ Принято с замечаниями", "⚠️ Ескертулермен қабылданды"),
            "needs_rework": T("❌ Требует доработки", "❌ Қайта қарауды қажет етеді"),
        }.get(v_raw, v_raw))
        assessment_html = f"""
        <div class="box ai-box">
            <h3>{T('🤖 ИИ-ЗАКЛЮЧЕНИЕ ЦИФРОВОГО КОНТРОЛЁРА («НарядAI»)', '🤖 ЦИФРЛЫҚ БАҚЫЛАУШЫНЫҢ ЖИ-ҚОРЫТЫНДЫСЫ («НарядAI»)')}</h3>
            <p><strong>{T('Вердикт ИИ:', 'ЖИ үкімі:')}</strong> {verdict_trans} | <strong>{T('Балл:', 'Балл:')}</strong> {_e(a.score)}/100
               {f'| <strong>{T('Мастер скорректировал оценку:', 'Шебер бағаны түзетті:')}</strong> {_e(a.master_score)}/100' if a.master_score is not None else ''}</p>
            <p><strong>{T('Пояснение:', 'Түсініктеме:')}</strong> {_e(a.explanation or '—')}</p>
            {f'<p><strong>{T('Комментарий мастера:', 'Шебердің түсініктемесі:')}</strong> {_e(a.master_comment)}</p>' if a.master_comment else ''}
        </div>
        """

    return f"""<!DOCTYPE html>
<html lang="{T('ru', 'kk')}">
<head>
<meta charset="utf-8">
<title>{T('Наряд-задание', 'Наряд-тапсырма')} №{order.number} — {T('АО «Костанайские Минералы»', '«Қостанай минералдары» АҚ')}</title>
<style>
    @page {{ size: A4; margin: 15mm; }}
    body {{
        font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Arial, sans-serif;
        font-size: 11pt;
        color: #1e293b;
        line-height: 1.4;
        margin: 0;
        padding: 10px;
    }}
    .header {{
        text-align: center;
        border-bottom: 2px solid #065f46;
        padding-bottom: 10px;
        margin-bottom: 15px;
    }}
    .header h1 {{ margin: 0; font-size: 16pt; color: #065f46; }}
    .header h2 {{ margin: 4px 0 0; font-size: 13pt; font-weight: normal; color: #334155; }}
    .meta-grid {{
        display: grid;
        grid-template-columns: 1fr 1fr;
        gap: 8px;
        margin-bottom: 15px;
    }}
    .box {{
        background: #f8fafc;
        border: 1px solid #cbd5e1;
        border-radius: 6px;
        padding: 8px 12px;
    }}
    .box h3 {{ margin: 0 0 6px; font-size: 11pt; color: #0f172a; border-bottom: 1px solid #e2e8f0; padding-bottom: 4px; }}
    .ai-box {{ background: #f0fdf4; border-color: #86efac; }}
    .ai-box h3 {{ color: #166534; }}
    table {{ width: 100%; border-collapse: collapse; margin-top: 6px; font-size: 9.5pt; }}
    th, td {{ border: 1px solid #cbd5e1; padding: 5px 8px; text-align: left; }}
    th {{ background: #f1f5f9; font-weight: bold; }}
    .signatures {{
        display: grid;
        grid-template-columns: 1fr 1fr 1fr;
        gap: 15px;
        margin-top: 30px;
        padding-top: 15px;
        border-top: 1px dashed #94a3b8;
    }}
    .sig-line {{ border-bottom: 1px solid #000; height: 30px; margin-top: 5px; }}
    .badge {{ display: inline-block; padding: 2px 6px; border-radius: 4px; font-size: 9pt; font-weight: bold; }}
    .badge-emergency {{ background: #fee2e2; color: #991b1b; }}
    .badge-normal {{ background: #e0f2fe; color: #075985; }}
    @media print {{
        body {{ padding: 0; }}
        .no-print {{ display: none; }}
    }}
</style>
</head>
<body>

<div class="no-print" style="margin-bottom: 15px; text-align: right;">
    <button onclick="window.print()" style="padding: 8px 16px; background: #059669; color: #fff; border: none; border-radius: 6px; font-weight: bold; cursor: pointer;">
        🖨️ {T('Распечатать / Сохранить в PDF', 'Басып шығару / PDF-ке сақтау')}
    </button>
</div>

<div class="header">
    <h1>{T('АО «КОСТАНАЙСКИЕ МИНЕРАЛЫ»', '«ҚОСТАНАЙ МИНЕРАЛДАРЫ» АҚ')}</h1>
    <h2>{T('НАРЯД-ЗАДАНИЕ №', 'ЖӨНДЕУ ЖҰМЫСТАРЫН ОРЫНДАУҒА НАРЯД-ТАПСЫРМА №')} {order.number}{T(' НА ВЫПОЛНЕНИЕ РЕМОНТНЫХ РАБОТ', '')}</h2>
    <div style="font-size: 9pt; color: #64748b; margin-top: 4px;">{T('Система автоматизированного контроля «НарядAI»', '«НарядAI» автоматтандырылған бақылау жүйесі')}</div>
</div>

<div class="meta-grid">
    <div class="box">
        <h3>{T('ОБОРУДОВАНИЕ И УЧАСТОК', 'ЖАБДЫҚ ЖӘНЕ БӨЛІМШЕ')}</h3>
        <div><strong>{T('Оборудование:', 'Жабдық:')}</strong> {eq_desc}</div>
        <div><strong>{T('Участок:', 'Бөлімше:')}</strong> {sec_name}</div>
        <div><strong>{T('Тип работ:', 'Жұмыс түрі:')}</strong> {T('Внеплановый (аварийный)', 'Жоспардан тыс (апаттық)') if order.work_type == 'unplanned' else T('Плановый регламентный', 'Жоспарлы регламенттік')}</div>
        <div><strong>{T('Приоритет:', 'Басымдық:')}</strong> <span class="badge {'badge-emergency' if order.priority == 'emergency' else 'badge-normal'}">{priority_label}</span></div>
    </div>
    <div class="box">
        <h3>{T('СРОКИ И ОТВЕТСТВЕННЫЕ', 'МЕРЗІМДЕР ЖӘНЕ ЖАУАПТЫЛАР')}</h3>
        <div><strong>{T('Выдан:', 'Берілді:')}</strong> {created_str}</div>
        <div><strong>{T('Срок исполнения:', 'Орындау мерзімі:')}</strong> {deadline_str}</div>
        <div><strong>{T('Мастер смены:', 'Ауысым шебері:')}</strong> {master_name}</div>
        <div><strong>{T('Исполнитель:', 'Орындаушы:')}</strong> {assignee_name} ({assignee_spec})</div>
    </div>
</div>

<div class="box" style="margin-bottom: 12px;">
    <h3>{T('ОПИСАНИЕ НЕИСПРАВНОСТИ / ЗАДАНИЕ МАСТЕРА', 'АҚАУДЫҢ СИПАТТАМАСЫ / ШЕБЕРДІҢ ТАПСЫРМАСЫ')}</h3>
    <p style="margin: 4px 0;">{description}</p>
    {f'<div style="font-size: 9pt; color: #64748b;">{T('Комментарий: {comment}', 'Түсініктеме: {comment}')}</div>' if comment else ''}
</div>

<div class="box" style="margin-bottom: 12px;">
    <h3>{T('ФАКТИЧЕСКИ ВЫПОЛНЕННЫЕ РАБОТЫ И МАТЕРИАЛЫ', 'НАҚТЫ ОРЫНДАЛҒАН ЖҰМЫСТАР МЕН МАТЕРИАЛДАР')}</h3>
    <div><strong>{T('Шифр неисправности:', 'Ақау шифры:')}</strong> [{fault_code}] {fault_name}</div>
    <div><strong>{T('Выполненные операции:', 'Орындалған операциялар:')}</strong> {work_done}</div>
    {f'<div><strong>{T('Комментарий исполнителя:', 'Орындаушының түсініктемесі:')}</strong> {close_comment}</div>' if close_comment else ''}
    
    <h4 style="margin: 8px 0 2px; font-size: 9.5pt;">{T('Списанные материалы и запчасти:', 'Есептен шығарылған материалдар мен қосалқы бөлшектер:')}</h4>
    <table>
        <thead><tr><th>{T('Материал / Запчасть', 'Материал / Қосалқы бөлшек')}</th><th style="width: 120px;">{T('Количество', 'Саны')}</th></tr></thead>
        <tbody>{mats_html}</tbody>
    </table>
</div>

{assessment_html}

<div class="box" style="margin-bottom: 15px;">
    <h3>{T('ХРОНОЛОГИЯ ПЕРЕХОДОВ СТАТУСОВ', 'МӘРТЕБЕ ӨТУЛЕРІНІҢ ХРОНОЛОГИЯСЫ')}</h3>
    <table>
        <thead><tr><th>{T('Время', 'Уақыт')}</th><th>{T('Действие', 'Әрекет')}</th><th>{T('Автор', 'Авторы')}</th><th>{T('Примечание / Причина', 'Ескертпе / Себеп')}</th></tr></thead>
        <tbody>{events_html}</tbody>
    </table>
</div>

<div class="signatures">
    <div>
        <strong>{T('Мастер смены:', 'Ауысым шебері:')}</strong>
        <div class="sig-line"></div>
        <div style="font-size: 8.5pt; color: #64748b; margin-top: 2px;">{master_name}</div>
    </div>
    <div>
        <strong>{T('Исполнитель:', 'Орындаушы:')}</strong>
        <div class="sig-line"></div>
        <div style="font-size: 8.5pt; color: #64748b; margin-top: 2px;">{assignee_name}</div>
    </div>
    <div>
        <strong>{T('Начальник участка / службы:', 'Бөлімше / қызмет бастығы:')}</strong>
        <div class="sig-line"></div>
        <div style="font-size: 8.5pt; color: #64748b; margin-top: 2px;">{T('Подпись / Дата', 'Қолы / Күні')}</div>
    </div>
</div>

</body>
</html>
"""
