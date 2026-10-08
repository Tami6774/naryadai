import io
import re
import sys
import unittest
from datetime import datetime, timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from openpyxl import load_workbook

import test_exports_nlp as base
from app.i18n import set_lang
from app.models import WorkOrder
from app.services.export import (
    export_materials_excel, export_rating_excel, export_shift_report_excel,
    generate_order_print_html,
)

KZ_LETTERS = re.compile(r"[әғқңөұүһі]", re.I)


def _texts(buf: io.BytesIO) -> list[tuple[str, str]]:
    wb = load_workbook(io.BytesIO(buf.getvalue()))
    out = []
    for ws in wb.worksheets:
        out.append(("sheet", ws.title))
        for row in ws.iter_rows(values_only=True):
            out.extend(("cell", str(v)) for v in row if isinstance(v, str))
    return out


class TestExportKazakh(unittest.TestCase):
    # Та же тестовая БД, что и у test_exports_nlp
    setUpClass = classmethod(base.TestExportsAndNLP.setUpClass.__func__)
    tearDownClass = classmethod(base.TestExportsAndNLP.tearDownClass.__func__)

    def tearDown(self):
        set_lang("ru")

    def _period(self):
        now = datetime.now()
        return now - timedelta(days=1), now + timedelta(days=1)

    def test_excel_headers_kazakh(self):
        start, end = self._period()
        with self.TestSession() as db:
            for fn in (export_shift_report_excel, export_rating_excel, export_materials_excel):
                set_lang("kz")
                texts = _texts(fn(db, start, end))
                self.assertTrue(any(KZ_LETTERS.search(t) for k, t in texts if k == "sheet"), fn.__name__)
                self.assertTrue(any(KZ_LETTERS.search(t) for k, t in texts), fn.__name__)
                for k, t in texts:
                    if k == "sheet":
                        self.assertLessEqual(len(t), 31)
                        self.assertFalse(set(t) & set("[]:*?/\\"))
                set_lang("ru")
                ru_sheets = [t for k, t in _texts(fn(db, start, end)) if k == "sheet"]
                self.assertFalse(any(KZ_LETTERS.search(t) for t in ru_sheets))

    def test_print_form_kazakh(self):
        with self.TestSession() as db:
            order = db.query(WorkOrder).filter_by(number=101).one()
            set_lang("kz")
            kz = generate_order_print_html(order)
            set_lang("ru")
            ru = generate_order_print_html(order)
        self.assertIn('<html lang="kk">', kz)
        self.assertIn("НАРЯД-ТАПСЫРМА", kz)
        self.assertIn("Апаттық", kz)
        self.assertNotIn("ОБОРУДОВАНИЕ И УЧАСТОК", kz)
        self.assertIn('<html lang="ru">', ru)
        self.assertIn("ОБОРУДОВАНИЕ И УЧАСТОК", ru)


if __name__ == "__main__":
    unittest.main()
